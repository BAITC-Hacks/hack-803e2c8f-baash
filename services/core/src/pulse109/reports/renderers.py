"""PDF/XLSX renderers consuming the exact AnalyticsResult object."""

from __future__ import annotations

import hashlib
import json
import textwrap
from io import BytesIO

from pulse109.analytics.models import AnalyticsResult


class RendererUnavailable(RuntimeError):
    pass


def _xlsx_value(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def render_pdf(result: AnalyticsResult, *, watermark: str = "SYNTHETIC / GOVERNED") -> bytes:
    try:
        from reportlab.lib import colors  # type: ignore[import-untyped]
        from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
        from reportlab.pdfgen import canvas  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RendererUnavailable("reportlab is required for PDF rendering") from exc
    output = BytesIO()
    document = canvas.Canvas(output, pagesize=A4)
    document.setTitle(f"Pulse 109 {result.metric_id}")

    width, height = A4

    def draw_header(page_number: int) -> float:
        document.setFillColor(colors.HexColor("#172421"))
        document.rect(0, height - 74, width, 74, fill=1, stroke=0)
        document.setFillColor(colors.white)
        document.setFont("Helvetica-Bold", 18)
        document.drawString(40, height - 42, "Pulse 109 situation report")
        document.setFont("Helvetica", 8)
        document.drawRightString(width - 40, height - 42, watermark)

        document.setFillColor(colors.HexColor("#12201D"))
        document.setFont("Helvetica-Bold", 11)
        document.drawString(
            40, height - 104, f"Metric: {result.metric_id} v{result.metric_version}"
        )
        document.setFont("Helvetica", 9)
        document.drawString(40, height - 121, f"Cutoff: {result.data_cutoff.isoformat()}")
        document.drawString(40, height - 136, f"Quality: {result.quality}")
        document.drawRightString(width - 40, height - 136, f"Page {page_number}")

        document.setFont("Helvetica", 8)
        document.drawString(
            40,
            height - 152,
            f"Counted: {result.records_considered}; "
            f"excluded untrusted time: {result.excluded_records}",
        )
        document.drawString(
            40,
            height - 165,
            f"Regions: {len(result.coverage)}; missing: {len(result.missing_regions)}",
        )

        header_y = height - 205
        document.setFillColor(colors.HexColor("#E8F1EE"))
        document.roundRect(40, header_y - 15, width - 80, 24, 3, fill=1, stroke=0)
        document.setFillColor(colors.HexColor("#12201D"))
        document.setFont("Helvetica-Bold", 9)
        document.drawString(48, header_y - 7, " | ".join(column.name for column in result.columns))
        return float(header_y - 35)

    page_number = 1
    y = draw_header(page_number)
    document.setFont("Helvetica", 9)
    for row in result.rows:
        document.setFillColor(colors.HexColor("#243531"))
        document.drawString(48, y, " | ".join(str(value) for value in row))
        document.setStrokeColor(colors.HexColor("#D7E0DD"))
        document.line(40, y - 7, width - 40, y - 7)
        y -= 22
        if y < 55:
            document.showPage()
            page_number += 1
            y = draw_header(page_number)
            document.setFont("Helvetica", 9)

    if not result.rows:
        document.setFillColor(colors.HexColor("#6C7B77"))
        document.drawString(48, y, "No numeric rows available for this governed cutoff.")
        y -= 22
    metadata_lines = [
        "Coverage: " + ", ".join(f"{region}={state}" for region, state in result.coverage.items()),
        *("Source: " + reference for reference in result.provenance),
    ]
    if result.executed_query is not None:
        query = result.executed_query
        metadata_lines.extend(
            [
                (
                    f"Period: {query.time_from.isoformat()} to "
                    f"{query.time_to.isoformat()} (end exclusive)"
                ),
                "Filters: "
                + json.dumps(
                    [item.model_dump(mode="json") for item in query.filters], ensure_ascii=True
                ),
                f"Granularity: {query.granularity}; buckets: UTC",
            ]
        )
    for line in metadata_lines:
        for fragment in textwrap.wrap(line, width=95):
            if y < 55:
                document.showPage()
                page_number += 1
                y = draw_header(page_number)
            document.setFillColor(colors.HexColor("#243531"))
            document.setFont("Helvetica", 8)
            document.drawString(40, y, fragment)
            y -= 13
    document.save()
    return output.getvalue()


def render_xlsx(result: AnalyticsResult, *, watermark: str = "SYNTHETIC / GOVERNED") -> bytes:
    try:
        from openpyxl import Workbook  # type: ignore[import-untyped]
        from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
        from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RendererUnavailable("openpyxl is required for XLSX rendering") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = result.metric_id[:31]
    sheet.append([f"Metric: {result.metric_id} v{result.metric_version}"])
    sheet.append([f"Data cutoff: {result.data_cutoff.isoformat()}", f"Watermark: {watermark}"])
    sheet.append([column.name for column in result.columns])
    for row in result.rows:
        sheet.append([_xlsx_value(value) for value in row])

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = (
        f"A3:{get_column_letter(max(1, len(result.columns)))}{max(3, sheet.max_row)}"
    )
    metadata = workbook.create_sheet("Provenance")
    metadata.append(["Metric", result.metric_id, result.metric_version])
    metadata.append(["Data cutoff", result.data_cutoff.isoformat()])
    metadata.append(["Computed at", result.computed_at.isoformat()])
    metadata.append(["Quality", result.quality])
    metadata.append(["Records considered", result.records_considered])
    metadata.append(["Excluded untrusted time", result.excluded_records])
    metadata.append(["Truncated", result.truncated])
    metadata.append(["Synthetic", result.synthetic])
    if result.executed_query is not None:
        query = result.executed_query
        metadata.append(["Period from", query.time_from.isoformat()])
        metadata.append(["Period to (exclusive)", query.time_to.isoformat()])
        metadata.append(["Granularity (UTC)", query.granularity])
        metadata.append(
            [
                "Filters",
                json.dumps(
                    [item.model_dump(mode="json") for item in query.filters], ensure_ascii=True
                ),
            ]
        )
    for region, state in result.coverage.items():
        metadata.append(["Region coverage", _xlsx_value(region), _xlsx_value(state)])
    for source in result.provenance:
        metadata.append(["Source", _xlsx_value(source)])
    metadata.column_dimensions["A"].width = 28
    metadata.column_dimensions["B"].width = 70
    sheet.row_dimensions[1].height = 24
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="172421")
    for cell in sheet[2]:
        cell.font = Font(color="51615D", italic=True)
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="12201D")
        cell.fill = PatternFill("solid", fgColor="E8F1EE")
        cell.alignment = Alignment(vertical="center")
    for column_index, column in enumerate(result.columns, start=1):
        values = [column.name, *(str(row[column_index - 1]) for row in result.rows)]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            48, max(14, max(len(value) for value in values) + 2)
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def artifact_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
