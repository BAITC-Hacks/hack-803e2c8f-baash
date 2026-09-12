from io import BytesIO
from uuid import UUID

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pulse109.main import app, report_runtime
from pypdf import PdfReader

QUERY = {
    "metric_id": "coverage",
    "dimensions": ["region_id"],
    "filters": {},
    "time_range": {
        "from": "2026-09-01T00:00:00Z",
        "to": "2026-09-12T00:00:00Z",
    },
    "granularity": "day",
}


def test_metric_and_both_exports_share_definition_cutoff_and_rows() -> None:
    client = TestClient(app)
    response = client.post("/v1/analytics/query", headers={"X-Region-Id": "ALL"}, json=QUERY)
    assert response.status_code == 200
    metric = response.json()
    assert metric["metric_id"] == "coverage"
    assert metric["metric_version"] == "1.0.0"
    assert metric["missing_regions"] == ["KAR"]
    assert all(row[0] != "KAR" for row in metric["rows"])

    jobs = []
    for report_format in ("pdf", "xlsx"):
        created = client.post(
            "/v1/reports",
            headers={
                "Idempotency-Key": f"m6-{report_format}-export-0001",
                "X-Region-Id": "ALL",
                "X-Actor-Token": "synthetic-analyst",
                "X-Export-Purpose": "M6 acceptance comparison",
            },
            json={
                "format": report_format,
                "template_id": "situation-center-1.0.0",
                "query": QUERY,
                "locale": "ru-KZ",
            },
        )
        assert created.status_code == 202
        assert created.json()["status"] == "succeeded"
        jobs.append(UUID(created.json()["job_id"]))

    pdf_artifact = report_runtime.artifacts[jobs[0]]
    xlsx_artifact = report_runtime.artifacts[jobs[1]]
    assert (pdf_artifact.metric_id, pdf_artifact.metric_version, pdf_artifact.data_cutoff) == (
        xlsx_artifact.metric_id,
        xlsx_artifact.metric_version,
        xlsx_artifact.data_cutoff,
    )
    assert pdf_artifact.data_cutoff.isoformat().replace("+00:00", "Z") == metric["data_cutoff"]

    pdf_text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(BytesIO(report_runtime.contents[jobs[0]])).pages
    )
    assert "Metric: coverage v1.0.0" in pdf_text
    assert "ALA | present" in pdf_text
    assert "KAR" not in pdf_text

    workbook = load_workbook(BytesIO(report_runtime.contents[jobs[1]]), data_only=False)
    sheet = workbook.active
    assert sheet["A1"].value == "Metric: coverage v1.0.0"
    rows = list(sheet.iter_rows(min_row=4, values_only=True))
    assert ("ALA", "present") in rows
    assert all(row[0] != "KAR" for row in rows)


def test_region_scope_cannot_be_bypassed_with_filter() -> None:
    client = TestClient(app)
    query = {**QUERY, "filters": {"region_id": ["AST"]}}
    response = client.post("/v1/analytics/query", headers={"X-Region-Id": "ALA"}, json=query)
    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "region_scope_denied"


def test_missing_source_alert_is_visible_not_numeric_zero() -> None:
    client = TestClient(app)
    response = client.get("/v1/alerts", headers={"X-Region-Id": "KAR"})
    assert response.status_code == 200
    alert = response.json()[0]
    assert alert["type"] == "data_quality"
    assert alert["evidence"]["state"] == "missing"
    assert alert["observed_value"] is None
